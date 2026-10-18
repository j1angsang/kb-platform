import logging
import os
import re
import threading
import time
from pathlib import Path

from services.storage import STORAGE_DIR, derived_dir_for

logger = logging.getLogger(__name__)

PDF_MIN_CHARS_PER_PAGE = float(os.getenv("PDF_MIN_CHARS_PER_PAGE", "10"))

PDF_OCR_MAX_PAGES = int(os.getenv("PDF_OCR_MAX_PAGES", "50"))

PDF_OCR_TEXT_ENABLED = (os.getenv("PDF_OCR_TEXT_ENABLED") or "false").strip().lower() in (
    "1", "true", "yes", "on",
)
PDF_OCR_TEXT_MAX_PAGES = int(os.getenv("PDF_OCR_TEXT_MAX_PAGES", "50"))
PDF_OCR_TEXT_MIN_CHARS = int(os.getenv("PDF_OCR_TEXT_MIN_CHARS", "5"))
PDF_OCR_TEXT_MIN_CONFIDENCE = float(os.getenv("PDF_OCR_TEXT_MIN_CONFIDENCE", "0.6"))
# 单页 OCR 硬超时，降级保留图片块，不许卡死入库管线
PDF_OCR_TEXT_TIMEOUT_SECONDS = float(os.getenv("PDF_OCR_TEXT_TIMEOUT_SECONDS", "120"))

CHUNK_SIZE = int(os.getenv("CHUNK_SIZE", "500"))
CHUNK_OVERLAP = int(os.getenv("CHUNK_OVERLAP", "50"))
if CHUNK_SIZE < 100:
    raise RuntimeError(f"CHUNK_SIZE 必须 >= 100，当前 {CHUNK_SIZE}")
if not 0 <= CHUNK_OVERLAP < CHUNK_SIZE:
    raise RuntimeError(
        f"CHUNK_OVERLAP 必须满足 0 <= overlap < size，当前 {CHUNK_OVERLAP}/{CHUNK_SIZE}"
    )


class PdfWithoutTextLayerError(ValueError):
    """PDF 抽不出可用文本层，疑似扫描件，供 indexer 记账"""


_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_SENTENCE_SPLIT_RE = re.compile(r"(?<=[。！？!?；;])")


def _split_sections(text: str) -> list[tuple[str, str]]:
    """按 Markdown 标题切成 (正文, 章节路径)，没有标题就整篇一段。"""
    sections: list[tuple[str, str]] = []
    stack: list[tuple[int, str]] = []
    buf: list[str] = []

    def flush() -> None:
        body = "\n".join(buf).strip()
        if body:
            sections.append((body, " > ".join(title for _, title in stack)))
        buf.clear()

    for line in text.splitlines():
        match = _HEADING_RE.match(line.strip())
        if match:
            flush()
            level = len(match.group(1))
            title = match.group(2).strip()
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, title))
        else:
            buf.append(line)
    flush()
    return sections or [(text.strip(), "")]


def _hard_split(text: str, size: int, overlap: int) -> list[str]:
    step = max(1, size - overlap)
    return [
        text[start:start + size]
        for start in range(0, len(text), step)
        if text[start:start + size].strip()
    ]


def _units_for_paragraph(paragraph: str, size: int) -> list[str]:
    """长段落优先按句子切；表格行按行切，避免把一行数据拆开。"""
    lines = [line.strip() for line in paragraph.splitlines() if line.strip()]
    table_like = len(lines) > 1 and sum(1 for line in lines if " | " in line) >= max(1, len(lines) // 2)
    if table_like:
        return lines
    if len(paragraph) <= size:
        return [paragraph]
    sentences = [part.strip() for part in _SENTENCE_SPLIT_RE.split(paragraph) if part.strip()]
    return sentences or [paragraph]


def _pack_units(units: list[str], size: int, overlap: int) -> list[str]:
    chunks: list[str] = []
    buf = ""
    for unit in units:
        unit = unit.strip()
        if not unit:
            continue
        if len(unit) > size:
            if buf:
                chunks.append(buf)
                buf = ""
            chunks.extend(_hard_split(unit, size, overlap))
            continue
        candidate = f"{buf}\n{unit}" if buf else unit
        if len(candidate) <= size:
            buf = candidate
            continue
        if buf:
            chunks.append(buf)
            tail = buf[-overlap:].strip() if overlap else ""
        else:
            tail = ""
        if tail and len(tail) + 1 + len(unit) <= size:
            buf = f"{tail}\n{unit}"
        else:
            buf = unit
    if buf:
        chunks.append(buf)
    return chunks


def _chunk_section(body: str, heading_path: str, size: int, overlap: int) -> list[str]:
    heading_path = (heading_path or "")[:80]
    prefix = f"[{heading_path}]\n" if heading_path else ""
    body_size = max(100, size - len(prefix))
    paragraphs = [part.strip() for part in re.split(r"\n\s*\n", body) if part.strip()]
    units: list[str] = []
    for paragraph in paragraphs:
        units.extend(_units_for_paragraph(paragraph, body_size))
    return [f"{prefix}{chunk}" for chunk in _pack_units(units, body_size, overlap)]


def _merge_chunks(chunks: list[str], size: int) -> list[str]:
    """把相邻小节合并到接近 size，避免“每个标题一个 chunk”切得过碎。"""
    merged: list[str] = []
    buf = ""
    for chunk in chunks:
        if not chunk.strip():
            continue
        candidate = f"{buf}\n\n{chunk}" if buf else chunk
        if len(candidate) <= size:
            buf = candidate
        else:
            if buf:
                merged.append(buf)
            buf = chunk
    if buf:
        merged.append(buf)
    return merged


def chunk_text(text: str, size: int | None = None, overlap: int | None = None) -> list[str]:
    """结构感知切分：先按标题分节，再按段落/句子/表格行打包。

    - 标题路径写入 chunk 前缀，检索命中后能直接看到「哪一章」；
    - 长段落优先在句号/问号/分号处断开，不再硬切字符；
    - 表格行尽量按行保留，避免一行数据被拆到两个 chunk；
    - 相邻小节会合并到接近 size，避免切得过碎。
    """
    chunk_size = CHUNK_SIZE if size is None else int(size)
    chunk_overlap = CHUNK_OVERLAP if overlap is None else int(overlap)
    if chunk_size < 100:
        raise ValueError(f"chunk size 必须 >= 100，当前 {chunk_size}")
    if not 0 <= chunk_overlap < chunk_size:
        raise ValueError(
            f"chunk overlap 必须满足 0 <= overlap < size，当前 {chunk_overlap}/{chunk_size}"
        )
    chunks: list[str] = []
    for body, heading_path in _split_sections(text):
        chunks.extend(_chunk_section(body, heading_path, chunk_size, chunk_overlap))
    return _merge_chunks([chunk for chunk in chunks if chunk.strip()], chunk_size)

def _read_text(path: Path) -> str:
    data = path.read_bytes()
    for enc in ("utf-8", "gbk"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")

IMAGE_SUFFIXES = (".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp")

# 上传白名单与解析能力必须同一份，否则会"提交成功、几秒后才报不支持"
SUPPORTED_UPLOAD_SUFFIXES = (".txt", ".md", ".pdf", ".docx", ".xlsx") + IMAGE_SUFFIXES


def is_supported_upload(filename: str) -> bool:
    return (Path(filename).suffix or "").lower() in SUPPORTED_UPLOAD_SUFFIXES


def parse_file(path: Path) -> list[dict]:
    started = time.perf_counter()
    suffix = path.suffix.lower()
    if suffix in (".txt", ".md"):
        chunks = _parse_text(path)
    elif suffix in IMAGE_SUFFIXES:
        chunks = [{"type": "image", "seq": 0, "image_path": str(path)}]
    elif suffix == ".pdf":
        chunks = _parse_pdf(path)
    elif suffix == ".docx":
        chunks = _parse_docx(path)
    elif suffix == ".xlsx":
        chunks = _parse_xlsx(path)
    else:
        raise ValueError(f"暂不支持解析:{suffix}")

    logger.info(
        "parser file=%s suffix=%s chunks=%d elapsed_ms=%.1f",
        path.name, suffix, len(chunks), (time.perf_counter() - started) * 1000,
    )
    for ch in chunks:
        preview = (ch.get("content") or ch.get("image_path") or "").replace("\n", " ")
        logger.debug(
            "parser.chunk file=%s seq=%s type=%s len=%d preview=%s",
            path.name, ch.get("seq"), ch.get("type"),
            len(ch.get("content") or ""), preview[:60],
        )
    return chunks


def _parse_text(path: Path) -> list[dict]:
    return [
        {"type": "text", "seq": i, "content": c}
        for i, c in enumerate(chunk_text(_read_text(path)))
    ]


def _parse_pdf(path: Path) -> list[dict]:
    from pypdf import PdfReader  # 延迟导入:不碰 pdf 就不加载这个库

    reader = PdfReader(str(path))
    out: list[dict] = []
    seq = 0
    total_chars = 0
    textless_pages: list[int] = []  # 扫描件整本如此，混合件部分页如此
    for page_no, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        total_chars += len(text.strip())
        if len(text.strip()) >= PDF_MIN_CHARS_PER_PAGE:
            for piece in chunk_text(text):
                out.append({"type": "text", "seq": seq, "content": piece, "page": page_no})
                seq += 1
        else:
            textless_pages.append(page_no)

    # 无文本层页渲染成图片块；渲染不可用就抛错，不能退回 ready+0 块的静默失效
    if textless_pages:
        renderable = textless_pages[:PDF_OCR_MAX_PAGES]
        if len(textless_pages) > PDF_OCR_MAX_PAGES:
            logger.warning(
                "parser.pdf_ocr_truncated file=%s pages=%d rendered=%d limit=%d 超上限截断，其余页未入库",
                path.name, len(textless_pages), len(renderable), PDF_OCR_MAX_PAGES,
            )
        ocr_pages: set[int] = set()
        if PDF_OCR_TEXT_ENABLED:
            ocr_pages = set(renderable[:PDF_OCR_TEXT_MAX_PAGES])
            if len(renderable) > PDF_OCR_TEXT_MAX_PAGES:
                logger.warning(
                    "parser.pdf_ocr_text_truncated file=%s pages=%d ocr=%d limit=%d "
                    "超出 OCR 页数上限的页只保留图片块",
                    path.name, len(renderable), len(ocr_pages), PDF_OCR_TEXT_MAX_PAGES,
                )
        try:
            for page_no in renderable:
                rendered = _render_pdf_page(path, page_no)
                out.append({
                    "type": "image",
                    "seq": seq,
                    "image_path": str(rendered),
                    "page": page_no,  # store.py 会把图片块的 page 写进 metadata
                })
                seq += 1
                if page_no not in ocr_pages:
                    continue
                try:
                    ocr_started = time.perf_counter()
                    ocr_text, confidence = _ocr_page_text_with_timeout(rendered)
                    elapsed_ms = (time.perf_counter() - ocr_started) * 1000
                    cleaned = _clean_ocr_text(ocr_text)
                    effective = len(re.sub(r"\W+", "", cleaned, flags=re.UNICODE))
                    if effective < PDF_OCR_TEXT_MIN_CHARS:
                        logger.warning(
                            "parser.pdf_ocr_text_skipped file=%s page=%d ocr_chars=%d "
                            "effective_chars=%d min_chars=%d engine=rapidocr reason=too_short",
                            path.name, page_no, len(cleaned), effective, PDF_OCR_TEXT_MIN_CHARS,
                        )
                        continue
                    source = "ocr"
                    if confidence < PDF_OCR_TEXT_MIN_CONFIDENCE:
                        source = "ocr_lowconf"
                        logger.warning(
                            "parser.pdf_ocr_text_low_confidence file=%s page=%d confidence=%.4f "
                            "threshold=%.2f engine=rapidocr 仍入库但打标 source=%s",
                            path.name, page_no, confidence, PDF_OCR_TEXT_MIN_CONFIDENCE, source,
                        )
                    logger.info(
                        "parser.pdf_ocr_text file=%s page=%d ocr_chars=%d elapsed_ms=%.1f "
                        "engine=rapidocr confidence=%.4f source=%s",
                        path.name, page_no, len(cleaned), elapsed_ms, confidence, source,
                    )
                    out.append({
                        "type": "text",
                        "seq": seq,
                        "content": cleaned,
                        "page": page_no,
                        "source": source,
                        "ocr_confidence": round(confidence, 4),
                    })
                    seq += 1
                except Exception as ocr_exc:
                    logger.warning(
                        "parser.pdf_ocr_text_failed file=%s page=%d engine=rapidocr "
                        "error=%s:%s 该页降级为仅图片块",
                        path.name, page_no, type(ocr_exc).__name__, ocr_exc,
                    )
        except Exception as exc:
            raise PdfWithoutTextLayerError(
                f"PDF 无文本层，疑似扫描件：文件={path.name} 页数={len(reader.pages)} "
                f"抽出字符={total_chars} 无文本层页={textless_pages} "
                f"页渲染失败={type(exc).__name__}: {exc} "
                f"→ 该文件需要 OCR/视觉解析能力，当前环境不可用"
            ) from exc
    return out


def _render_pdf_page(pdf_path: Path, page_no: int, scale: float = 2.0) -> Path:
    import pypdfium2 as pdfium  # 延迟导入：不碰渲染就不加载

    page_dir = derived_dir_for(pdf_path.name)
    page_dir.mkdir(parents=True, exist_ok=True)
    out_path = page_dir / f"page_{page_no:03d}.jpg"
    pdf = pdfium.PdfDocument(str(pdf_path))
    try:
        bitmap = pdf[page_no - 1].render(scale=scale)
        bitmap.to_pil().convert("RGB").save(out_path, format="JPEG", quality=85)
    finally:
        pdf.close()
    logger.info(
        "parser.pdf_page_rendered file=%s page=%d out=%s scale=%.1f",
        pdf_path.name, page_no, out_path.name, scale,
    )
    return out_path


_ocr_engine = None
_ocr_engine_lock = threading.Lock()


def _get_ocr_engine():
    global _ocr_engine
    if _ocr_engine is None:
        with _ocr_engine_lock:
            if _ocr_engine is None:
                from rapidocr import RapidOCR  # 延迟导入

                _ocr_engine = RapidOCR()
    return _ocr_engine


def _ocr_page_text(image_path: Path) -> tuple[str, float]:
    engine = _get_ocr_engine()
    result = engine(str(image_path))
    txts = getattr(result, "txts", None)
    if not txts:
        return "", 0.0
    scores = [float(s) for s in (getattr(result, "scores", None) or [])]
    confidence = sum(scores) / len(scores) if scores else 0.0
    return "\n".join(str(t) for t in txts), confidence


def _ocr_page_text_with_timeout(image_path: Path) -> tuple[str, float]:
    outcome: dict = {}

    def _run() -> None:
        try:
            outcome["value"] = _ocr_page_text(image_path)
        except Exception as exc:  # 线程内异常搬运回调用方
            outcome["error"] = exc

    worker = threading.Thread(
        target=_run, daemon=True, name=f"kb-ocr-page-{image_path.name}"
    )
    worker.start()
    worker.join(PDF_OCR_TEXT_TIMEOUT_SECONDS)
    if worker.is_alive():
        raise TimeoutError(
            f"OCR 单页超过 {PDF_OCR_TEXT_TIMEOUT_SECONDS}s 未完成"
        )
    if "error" in outcome:
        raise outcome["error"]
    return outcome["value"]


def _clean_ocr_text(raw: str) -> str:
    return "\n".join(line.strip() for line in (raw or "").splitlines() if line.strip())


def _parse_docx(path: Path) -> list[dict]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    def heading_prefix(paragraph: Paragraph) -> str:
        style_name = str(getattr(getattr(paragraph, "style", None), "name", "") or "")
        if "heading" not in style_name.lower() and "标题" not in style_name:
            return ""
        digits = re.findall(r"\d+", style_name)
        level = int(digits[0]) if digits else 1
        return "#" * max(1, min(level, 6)) + " "

    doc = Document(str(path))
    parts: list[str] = []
    # 必须按 body 顺序遍历，否则段落/表格交错时阅读顺序全乱
    for child in doc.element.body.iterchildren():
        tag = getattr(child, "tag", "")
        if tag.endswith("}p"):
            para = Paragraph(child, doc)
            text = para.text.strip()
            if text:
                parts.append(heading_prefix(para) + text)
        elif tag.endswith("}tbl"):
            table = Table(child, doc)
            for row in table.rows:
                row_text = " | ".join(c.text.strip() for c in row.cells)
                if row_text.strip(" |"):
                    parts.append(row_text)
    out: list[dict] = []
    for i, piece in enumerate(chunk_text("\n\n".join(parts))):
        out.append({"type": "text", "seq": i, "content": piece})
    return out


def _parse_xlsx(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    out: list[dict] = []
    seq = 0
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            vals = ["" if v is None else str(v) for v in row]
            if any(vals):
                rows.append(" | ".join(vals))
        text = "\n".join(rows)
        if not text.strip():
            continue
        for piece in chunk_text(text):
            out.append({"type": "text", "seq": seq, "content": f"[{ws.title}]\n{piece}", "sheet": ws.title})
            seq += 1
    wb.close()
    return out

if __name__ == "__main__":
    import sys

    from config.logging_conf import setup_logging

    setup_logging()
    for ch in parse_file(Path(sys.argv[1])):
        preview = (ch.get("content") or ch.get("image_path") or "").replace("\n", " ")
        logger.info(
            "parser.preview seq=%s type=%s len=%d preview=%s",
            ch.get("seq"), ch.get("type"), len(ch.get("content") or ""), preview[:60],
        )

